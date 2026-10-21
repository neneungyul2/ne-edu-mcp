"""학교알리미 공식 엑셀(OpenAPI_Output.xlsx, sido_sggCode.xlsx)에서 패키지 데이터 JSON을 만든다.

사용: python scripts/build_schoolinfo_data.py <OpenAPI_Output.xlsx> <sido_sggCode.xlsx>
엑셀은 https://www.schoolinfo.go.kr/download/ 에서 받는다. 시트 이름 끝의 (번호)가 apiType이다.
"""
import json
import re
import sys
from pathlib import Path

import openpyxl

OUT = Path(__file__).resolve().parent.parent / "src" / "ne_edu_mcp" / "data"


def build_items(path: str) -> list[dict]:
    items = []
    wb = openpyxl.load_workbook(path, read_only=True)
    for sheet in wb.sheetnames:
        m = re.fullmatch(r"(.+)\((\d+)\)", sheet.strip())
        if not m:
            raise SystemExit(f"시트 이름에서 apiType을 읽을 수 없습니다: {sheet!r}")
        rows = list(wb[sheet].iter_rows(values_only=True))
        header = rows[0]
        section_col = 4 if len(header) > 4 and header[4] else None
        fields = []
        for r in rows[1:]:
            if r[2] is None:
                continue
            f = {"id": r[2], "label": r[1], "levels": r[3]}
            if section_col is not None and r[section_col]:
                f["section"] = r[section_col]
            fields.append(f)
        items.append({"api_type": m.group(2), "name": m.group(1), "fields": fields})
    return items


def build_regions(path: str) -> list[dict]:
    ws = openpyxl.load_workbook(path, read_only=True).active
    sidos: dict[str, dict] = {}
    for sido, sido_code, sgg, sgg_code in ws.iter_rows(min_row=4, values_only=True):
        if not sido or sido == "전체":
            continue
        entry = sidos.setdefault(str(sido_code), {"code": str(sido_code), "name": sido, "sgg": []})
        entry["sgg"].append({"code": str(sgg_code), "name": sgg})
    return list(sidos.values())


if __name__ == "__main__":
    items_xlsx, regions_xlsx = sys.argv[1:3]
    OUT.mkdir(parents=True, exist_ok=True)
    (OUT / "schoolinfo_items.json").write_text(
        json.dumps(build_items(items_xlsx), ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    (OUT / "schoolinfo_regions.json").write_text(
        json.dumps(build_regions(regions_xlsx), ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    print("ok")
